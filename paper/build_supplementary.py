"""Builds the supplementary figures PDF: all four figures at full size with
their captions, plus a one-line header naming the paper.

Usage:
    python3 paper/build_supplementary.py
Then convert to PDF with LibreOffice headless:
    soffice --headless --convert-to pdf --outdir paper paper/supplementary_figures.docx
"""

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).parent
REPO_ROOT = HERE.parent
FIGURES_DIR = REPO_ROOT / "pilot" / "out" / "figures"
OUTPUT_DOCX = HERE / "supplementary_figures.docx"

PAPER_TITLE = "Anchored: Reference Exemplars Overwrite Language Model Self-Report"

FIGURES = [
    (
        "figure1_validity_screen.png",
        "Figure 1.",
        "Validity screen across five models and three response formats. A cell is Invalid when both anchors are constant across its 60 observations, in which case the rescaling is provably vacuous. The coarse format invalidates one model and the fine format invalidates a different one. The intermediate format invalidates none.",
    ),
    (
        "figure2_self_assessment_p5.png",
        "Figure 2.",
        "Mean self-assessment on p5 under three conditions. N carries no vignettes. V presents the vignettes before the self-question. R presents them after. Error bars are the standard error over 60 tasks. In every model the value under R returns to the value under N, which identifies turn order as the cause of the shift under V.",
    ),
    (
        "figure3_self_other_gap.png",
        "Figure 3.",
        "Signed self minus other gap on byte-identical code, normalised by scale width, for five models in three formats. The dashed line marks the preregistered failure threshold of 0.1875·W. Both Google models exceed the Anthropic models in every format, and the gap widens as the scale becomes finer.",
    ),
    (
        "figure4_confidence_accuracy_correlation.png",
        "Figure 4.",
        "Pearson correlation between stated confidence and execution-verified accuracy across five models, raw against rescaled, in three formats and in condition R. Accuracy excludes tasks where no code could be extracted. With five models no individual value is distinguishable from zero; the figure shows direction rather than effect size. Spearman coefficients are in the repository.",
    ),
]


def build() -> Path:
    document = docx.Document()

    header = document.add_paragraph()
    run = header.add_run(f"Supplementary Figures — {PAPER_TITLE}")
    run.bold = True
    run.font.size = Pt(12)

    for i, (filename, label, caption) in enumerate(FIGURES):
        document.add_picture(str(FIGURES_DIR / filename), width=Inches(6.3))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        p = document.add_paragraph()
        p.add_run(f"{label} ").bold = True
        p.add_run(caption)

        if i < len(FIGURES) - 1:
            document.add_page_break()

    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")

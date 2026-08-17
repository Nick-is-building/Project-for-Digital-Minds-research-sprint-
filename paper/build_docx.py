"""Regenerates the submission docx from paper_draft.md, using the official
template only for its page setup and styles (Title/Heading 2/Heading 3/normal).
The template's own placeholder content, guidance text and info box are never
copied — the body is cleared and rebuilt entirely from the markdown.

Usage:
    python3 paper/build_docx.py
Then convert to PDF with LibreOffice headless:
    soffice --headless --convert-to pdf --outdir paper paper/submission.docx
"""

import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

HERE = Path(__file__).parent
REPO_ROOT = HERE.parent
TEMPLATE = REPO_ROOT / "Digital Minds Research Sprint submission template.docx"
MARKDOWN = HERE / "paper_draft.md"
FIGURES_DIR = REPO_ROOT / "pilot" / "out" / "figures"
OUTPUT_DOCX = HERE / "submission.docx"

FIGURE_FILES = {
    1: "figure1_validity_screen.png",
    2: "figure2_self_assessment_p5.png",
    3: "figure3_self_other_gap.png",
    4: "figure4_confidence_accuracy_correlation.png",
}

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def _add_inline_runs(paragraph, text: str) -> None:
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)


def _clear_body(document) -> None:
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _add_page_numbers(document) -> None:
    footer = document.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = docx.oxml.OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = docx.oxml.OxmlElement("w:instrText")
    instr.text = "PAGE"
    fld_sep = docx.oxml.OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = docx.oxml.OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    for el in (fld_begin, instr, fld_sep, fld_end):
        r.append(el)


def _set_cell_borders(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = docx.oxml.OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = docx.oxml.OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tcPr.append(borders)


def _render_table(document, lines: list[str]) -> None:
    rows = [
        [c.strip() for c in line.strip().strip("|").split("|")]
        for line in lines
        if not all(re.match(r"^:?-{2,}:?$", c.strip()) for c in line.strip().strip("|").split("|"))
    ]
    n_cols = len(rows[0])
    table = document.add_table(rows=len(rows), cols=n_cols)
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = table.cell(ri, ci)
            _set_cell_borders(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(cell_text)
            if ri == 0:
                run.bold = True


FIGURE_CAPTION_RE = re.compile(r"^\*\*Figure (\d)\.\*\*")


def build() -> Path:
    document = docx.Document(TEMPLATE)
    _clear_body(document)
    _add_page_numbers(document)

    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        if line.startswith("| "):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            _render_table(document, table_lines)
            continue

        if line.startswith("#### "):
            document.add_heading(line[5:].strip(), level=4)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            p = document.add_paragraph(style="Title")
            _add_inline_runs(p, line[2:].strip())
        elif line.startswith("- "):
            p = document.add_paragraph(style="normal")
            _add_inline_runs(p, "• " + line[2:].strip())
        else:
            m = FIGURE_CAPTION_RE.match(line.strip())
            if m:
                fig_num = int(m.group(1))
                img_path = FIGURES_DIR / FIGURE_FILES[fig_num]
                document.add_picture(str(img_path), width=Inches(5.2))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = document.add_paragraph(style="normal")
            _add_inline_runs(p, line.strip())

        i += 1

    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
